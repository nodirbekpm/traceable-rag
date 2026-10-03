import io
import json
from contextlib import nullcontext
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import select
from sqlalchemy.orm import Session

from anchor import analysis
from anchor.api import main
from anchor.cache import MemoryCache
from anchor.config import Settings
from anchor.db import get_session
from anchor.documents import UploadError, media_type_for
from anchor.estimate import estimate_analysis
from anchor.extraction.extractor import windows
from anchor.models import ExtractedFact, ExtractionRun
from anchor.rerank import NoReranker
from anchor.storage import RawStore
from anchor.text import document_to_text, docx_to_text, pdf_to_text
from embed_fakes import FakeEmbedder
from llm_fakes import FakeLLM, FakeStreamingLLM

CONTRACT = (
    "SERVICE AGREEMENT\n\n"
    "This agreement is made on March 3, 2026 between Northwind Traders Ltd. and Contoso LLC.\n\n"
    "Contoso LLC will pay a fee of $120,000 per year for an initial term of 3 years.\n\n"
    "Either party may terminate with 90 days notice. Late payments accrue interest of 1.5% per month."
)


def make_pdf(text: str) -> bytes:
    """A one-page PDF with a real text layer, built by hand so tests need no PDF writer."""
    content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF" % (len(objects) + 1, xref)
    return out


def make_docx(paragraphs: list[str], table: list[list[str]]) -> bytes:
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    grid = document.add_table(rows=len(table), cols=len(table[0]))
    for r, row in enumerate(table):
        for c, value in enumerate(row):
            grid.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# --- formats ---------------------------------------------------------------------


def test_pdf_text_layer_is_extracted() -> None:
    assert "Total fee is 120000 USD" in pdf_to_text(make_pdf("Total fee is 120000 USD"))


def test_docx_paragraphs_and_tables_are_extracted() -> None:
    raw = make_docx(["Lease agreement", "Rent is $2,000 per month."], [["Deposit", "$4,000"]])

    assert docx_to_text(raw) == "Lease agreement\n\nRent is $2,000 per month.\n\nDeposit $4,000"


def test_plain_text_and_markdown_keep_their_content() -> None:
    assert document_to_text(b"# Title\r\n\r\nBody  text", "text/markdown") == "# Title\n\nBody text"


@pytest.mark.parametrize(
    ("name", "declared", "expected"),
    [
        ("report.PDF", None, "application/pdf"),
        ("notes.md", "application/octet-stream", "text/markdown"),
        ("page.htm", None, "text/html"),
        ("blob", "text/plain", "text/plain"),
    ],
)
def test_media_type_prefers_the_extension(name: str, declared: str | None, expected: str) -> None:
    assert media_type_for(name, declared) == expected


def test_unsupported_file_type_is_refused() -> None:
    with pytest.raises(UploadError, match="Unsupported file type"):
        media_type_for("photo.png", "image/png")


# --- long documents and the estimate ---------------------------------------------


def test_windows_cover_the_text_and_break_at_paragraphs() -> None:
    text = "\n\n".join(f"Paragraph {i} " + "x" * 300 for i in range(40))

    pieces = windows(text, size=2000)

    assert "".join(pieces) == text
    assert all(len(piece) <= 2000 for piece in pieces)
    assert all(piece.startswith("\n\nParagraph") for piece in pieces[1:])


def test_estimate_counts_calls_chunks_cost_and_rate_limit_waits(session: Session) -> None:
    settings = Settings(llm_requests_per_minute=2)
    long_text = "\n\n".join("Clause text " * 400 for _ in range(60))

    short = estimate_analysis(session, CONTRACT, settings)
    long = estimate_analysis(session, long_text, settings)

    assert (short.model_calls, short.pages, short.truncated) == (1, 1, False)
    assert short.chunks >= 1 and short.seconds > 0 and short.cost_usd > 0
    assert short.basis == "default rates (no history yet)"
    assert long.model_calls > 2
    assert long.seconds > short.seconds + 60
    assert any("rate limit" in note for note in long.notes)


# --- upload → estimate → analyse → ask ---------------------------------------------

GENERIC_FACTS = [
    {"field": "document_date", "value": "March 3, 2026", "quote": "made on March 3, 2026",
     "confidence": 0.95},
    {"field": "monetary_amount", "value": "$120,000", "entity": "annual fee",
     "quote": "a fee of $120,000 per year", "confidence": 0.9},
    {"field": "party", "value": "Northwind Traders Ltd.",
     "quote": "between Northwind Traders Ltd. and Contoso LLC", "confidence": 0.9},
]  # fmt: skip


@pytest.fixture
def client(session: Session, tmp_path: Path, monkeypatch):
    settings = Settings(raw_storage_dir=tmp_path)
    monkeypatch.setattr(main, "get_settings", lambda: settings)
    main.app.dependency_overrides[get_session] = lambda: session
    main.app.dependency_overrides[main.get_session_factory] = lambda: lambda: nullcontext(session)
    main.app.dependency_overrides[main.get_analysis_components] = lambda: {
        "llm": FakeLLM(GENERIC_FACTS),
        "embedder": FakeEmbedder(),
        "store": RawStore(tmp_path),
        "settings": settings,
    }
    analysis._jobs.clear()
    with TestClient(main.app) as client:
        yield client
    main.app.dependency_overrides.clear()


def upload(client, name: str, content: bytes) -> dict:
    response = client.post("/documents", files={"file": (name, content)})
    assert response.status_code == 201, response.text
    return response.json()


def test_upload_returns_an_estimate_and_does_not_analyse_yet(client, session) -> None:
    body = upload(client, "contract.txt", CONTRACT.encode())

    assert body["document"]["title"] == "contract.txt"
    assert body["document"]["doc_type"] == "upload"
    assert body["estimate"]["model_calls"] == 1
    assert body["status"]["state"] == "new"
    assert session.scalar(select(ExtractionRun)) is None


def test_same_file_twice_is_recognised(client) -> None:
    first = upload(client, "contract.txt", CONTRACT.encode())
    second = upload(client, "copy.txt", CONTRACT.encode())

    assert second["duplicate"] is True
    assert second["document"]["id"] == first["document"]["id"]


def test_bad_uploads_are_explained(client) -> None:
    assert client.post("/documents", files={"file": ("x.png", b"\x89PNG")}).status_code == 422
    empty = client.post("/documents", files={"file": ("empty.txt", b"")})
    assert empty.status_code == 422
    assert "empty" in empty.json()["detail"]


def test_analysis_extracts_generic_facts_with_spans_and_indexes(client, session) -> None:
    document_id = upload(client, "contract.txt", CONTRACT.encode())["document"]["id"]

    started = client.post(f"/documents/{document_id}/analyze")
    final = client.get(f"/documents/{document_id}/status").json()

    assert started.status_code == 202
    assert final["state"] == "ready", final
    run = session.scalar(select(ExtractionRun))
    assert run.schema_version == "g1"
    facts = {f.field_name: f for f in session.scalars(select(ExtractedFact))}
    assert facts["monetary_amount"].value_normalized == "120000"
    assert facts["monetary_amount"].entity_id == "annual fee"
    assert all(f.validation_status == "verified" for f in facts.values())
    text = client.get(f"/documents/{document_id}/text").json()["text"]
    span = facts["party"]
    assert text[span.span_start : span.span_end] == "Northwind Traders Ltd."


def test_questions_can_be_limited_to_one_document(client, session, tmp_path) -> None:
    first = upload(client, "contract.txt", CONTRACT.encode())["document"]["id"]
    other = upload(client, "other.txt", b"An unrelated memo about the annual fee of $5.")
    client.post(f"/documents/{first}/analyze")
    client.post(f"/documents/{other['document']['id']}/analyze")
    line = {"claim": "The fee is $120,000 per year.", "source": "S1",
            "quote": "a fee of $120,000 per year"}  # fmt: skip
    main.app.dependency_overrides[main.get_ask_components] = lambda: {
        "embedder": FakeEmbedder(),
        "reranker": NoReranker(),
        "llm": FakeStreamingLLM(json.dumps(line) + "\n"),
        "cache": MemoryCache(),
        "settings": Settings(raw_storage_dir=tmp_path),
    }

    response = client.get(
        "/ask", params={"q": "What is the annual fee?", "document_id": first, "mode": "text"}
    )

    sources = next(
        json.loads(row[6:])["sources"]
        for row in response.text.splitlines()
        if row.startswith("data: ") and '"sources"' in row
    )
    assert sources and {s["document_id"] for s in sources} == {first}
    assert "event: claim" in response.text
